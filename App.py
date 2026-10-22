import os
import io
import pandas as pd
import numpy as np
import streamlit as st
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

st.set_page_config(
    page_title="STAZE Blinkit Summary & Master Dashboard",
    page_icon="🔵",
    layout="wide"
)

st.markdown("""
<style>
    :root {
        --dash-blue: #4A90C2;
        --dash-blue-dark: #285B7A;
        --dash-blue-light: #EEF6FC;
        --dash-blue-soft: #DCECF8;
        --dash-border: #C8DCEB;
        --dash-text: #29465B;
        --lavender-light: #F3EBFB;
        --lavender-blue-mix: linear-gradient(135deg, #F3EBFB 0%, #EEF6FC 100%);
        --lavender-border: #D5C6E8;
    }
    .stApp {
        background: linear-gradient(180deg, #F6FAFD 0%, #FFFFFF 38%);
        color: var(--dash-text);
    }
    [data-testid="stHeader"] {
        background: rgba(255,255,255,0.92);
    }
    h1, h2, h3 {
        color: var(--dash-blue-dark) !important;
        font-weight: 700 !important;
        text-align: center !important;
    }
    [data-testid="stMetric"] {
        background: linear-gradient(135deg, #FFFFFF 0%, #EEF6FF 100%);
        border: 1px solid var(--dash-border);
        border-left: 5px solid #6AA6CF;
        border-radius: 14px;
        padding: 14px 18px;
        box-shadow: 0 3px 12px rgba(74, 144, 194, 0.10);
        text-align: center !important;
        align-items: center !important;
        justify-content: center !important;
    }
    [data-testid="stMetricLabel"] {
        color: #52708D !important;
        font-weight: 600 !important;
        text-align: center !important;
        display: flex;
        justify-content: center;
    }
    [data-testid="stMetricValue"] {
        color: var(--dash-blue-dark) !important;
        font-weight: 750 !important;
        text-align: center !important;
        display: flex;
        justify-content: center;
    }
    .stTabs [data-baseweb="tab-list"] {
        gap: 6px;
        background: var(--lavender-blue-mix);
        border: 1px solid var(--lavender-border);
        border-radius: 16px;
        padding: 6px;
        justify-content: center;
    }
    .stTabs [data-baseweb="tab"] {
        height: 42px;
        border-radius: 12px;
        color: #4A3B6C;
        font-weight: 600;
        background: #FFFFFF;
        border: 1px solid #E2D6F3;
        text-align: center;
    }
    .stTabs [aria-selected="true"] {
        background: linear-gradient(135deg, #6AA6CF, #4A90C2) !important;
        color: #FFFFFF !important;
        box-shadow: 0 3px 9px rgba(74, 144, 194, 0.25);
        border: 0 !important;
    }
    [data-baseweb="select"] > div {
        background: var(--lavender-blue-mix) !important;
        border-color: var(--lavender-border) !important;
        border-radius: 12px !important;
    }
    [data-testid="stFileUploader"] {
        background: #FFFFFF;
        border: 1px solid var(--dash-border);
        border-radius: 14px;
        padding: 8px;
        box-shadow: 0 2px 8px rgba(74, 144, 194, 0.07);
    }
    [data-testid="stFileUploaderDropzone"] {
        background: #F7FBFE;
        border: 1px dashed #9FC4DE;
        border-radius: 10px;
    }
    [data-testid="stFileUploader"] button,
    .stDownloadButton button,
    .stButton button {
        background: linear-gradient(135deg, #6AA6CF, #4A90C2) !important;
        color: #FFFFFF !important;
        border: 0 !important;
        border-radius: 12px !important;
        font-weight: 650 !important;
        box-shadow: 0 2px 7px rgba(11, 92, 173, 0.20);
    }
    [data-testid="stFileUploader"] button:hover,
    .stDownloadButton button:hover,
    .stButton button:hover {
        background: #397BAA !important;
        color: #FFFFFF !important;
    }
    [data-testid="stDataFrame"] {
        border: 1px solid var(--dash-border);
        border-radius: 12px;
        overflow: hidden;
        box-shadow: 0 2px 9px rgba(11, 92, 173, 0.06);
    }
    [data-testid="stDataFrame"] th {
        background: #E5F1F9 !important;
        color: #285B7A !important;
        font-weight: 700 !important;
        text-align: center !important;
    }
    hr {
        border-color: #D7E4F2 !important;
    }
</style>
""", unsafe_allow_html=True)

st.markdown("<h1 style='text-align: center;'>STAZE Blinkit Summary & Master Dashboard</h1>", unsafe_allow_html=True)
st.markdown("<p style='text-align: center;'>Upload Ad Campaign reports and Overall Sales reports. System automatically filters campaigns and products containing <b>'staze'</b>[cite: 3].</p>", unsafe_allow_html=True)

OVERALL_FRANCHISE_MASTER = {
    "Staze Gloss Lock 2-in-1 Liquid Lipstick (I'm Shy) 7 g": "Gloss Lock",
    "Staze Gloss Lock 2 in 1 Liquid Lipstick (I'm Pretty) 7 g": "Gloss Lock",
    "Staze Gloss Lock 2 in 1 Liquid Lipstick (I'm Happy) 7 g": "Gloss Lock",
    "Staze Gloss Lock 2 in 1 Liquid Lipstick (I'm Bubbly) 7 g": "Gloss Lock",
    "Staze Twin Flame 2 in 1 Lipstick + Lip Liner (03 Match-Made) 2 g": "Twin Flame",
    "Staze Twin Flame 2 in 1 Lipstick + Lip Liner (05 Hot Date) 2 g": "Twin Flame",
    "Staze 9to9 Y.U.M.M.Y. Lip Gloss Balm (04 Strawberry Jelly) 2 g": "Yummy",
    "Staze 9to9 Pout On Lip Crayon (10 Royally Ruby) 3.5 g": "Lip Crayon",
    "Staze 9to9 Pout On Lip Crayon (01 Nude Whisper) 3.5 g": "Lip Crayon",
    "Staze Lengthening & Volumizing Tubing Mascara (01 Blackest Black) 5 ml": "Mascara",
}

CITY_REGION_MASTER = {
    "Ahmedabad": "T-10", "Bengaluru": "T-10", "Chennai": "T-10", "Delhi": "T-10",
    "Gurgaon": "T-10", "Hyderabad": "T-10", "Kolkata": "T-10", "Lucknow": "T-10",
    "Mumbai": "T-10", "Pune": "T-10"
}

def get_region_for_city(city_val):
    if pd.isna(city_val): return "ROI"
    return CITY_REGION_MASTER.get(str(city_val).strip().title(), "ROI")

MONTH_ORDER_ASC = {
    "JANUARY": 1, "FEBRUARY": 2, "MARCH": 3, "APRIL": 4,
    "MAY": 5, "JUNE": 6, "JULY": 7, "AUGUST": 8,
    "SEPTEMBER": 9, "OCTOBER": 10, "NOVEMBER": 11, "DECEMBER": 12
}

def get_calendar_months_asc(df_input):
    if df_input is None or 'Month' not in df_input.columns: return []
    return sorted([str(x).strip().upper() for x in df_input['Month'].dropna().unique()], key=lambda x: MONTH_ORDER_ASC.get(x, 99))

def style_and_export_pivot(pivot_df, sheet_name="Comparison"):
    buffer = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    
    top_header_fill = PatternFill(start_color="4A90C2", end_color="4A90C2", fill_type="solid")
    top_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sec_header_fill = PatternFill(start_color="E5F1F9", end_color="E5F1F9", fill_type="solid")
    sec_header_font = Font(name="Calibri", size=10, bold=True, color="285B7A")
    index_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    index_font = Font(name="Calibri", size=10, bold=True, color="000000")
    total_fill = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")
    total_font = Font(name="Calibri", size=10, bold=True, color="000000")
    data_font = Font(name="Calibri", size=10, color="000000")
    thin_border = Border(left=Side(style='thin', color='BFBFBF'), right=Side(style='thin', color='BFBFBF'), top=Side(style='thin', color='BFBFBF'), bottom=Side(style='thin', color='BFBFBF'))

    if isinstance(pivot_df.columns, pd.MultiIndex):
        metrics = pivot_df.columns.get_level_values(1)
        entity_title = pivot_df.index.name if pivot_df.index.name else ""
        ws.cell(row=1, column=1, value="").fill = top_header_fill
        current_col = 2
        for month in pivot_df.columns.get_level_values(0).unique():
            sub_cols = [c for c in pivot_df.columns if c[0] == month]
            num_sub = len(sub_cols)
            if num_sub > 1: ws.merge_cells(start_row=1, start_column=current_col, end_row=1, end_column=current_col + num_sub - 1)
            cell = ws.cell(row=1, column=current_col, value=str(month))
            cell.fill = top_header_fill; cell.font = top_header_font; cell.alignment = Alignment(horizontal="center", vertical="center")
            for c in range(current_col, current_col + num_sub): ws.cell(row=1, column=c).border = thin_border
            current_col += num_sub
        ws.cell(row=2, column=1, value=str(entity_title)).fill = index_fill
        for col_idx, metric in enumerate(metrics, start=2):
            cell = ws.cell(row=2, column=col_idx, value=str(metric))
            cell.fill = sec_header_fill; cell.font = sec_header_font; cell.alignment = Alignment(horizontal="center", vertical="center")
        for r_idx, (idx_val, row_data) in enumerate(pivot_df.iterrows(), start=3):
            is_grand_total = (str(idx_val).strip().lower() in ['grand total', 'aggregate / average'])
            ws.cell(row=r_idx, column=1, value=str(idx_val)).fill = total_fill if is_grand_total else index_fill
            for c_idx, (m_col, val) in enumerate(zip(metrics, row_data), start=2):
                val_cell = ws.cell(row=r_idx, column=c_idx, value=val)
                if isinstance(val, (int, float, np.number)):
                    if str(m_col).upper() in ['ACOS', 'TACOS']:
                        val_cell.value = val / 100.0 if val > 1 else val
                        val_cell.number_format = '0.00%'
                    else:
                        val_cell.value = round(val, 2)
                        val_cell.number_format = '#,##0' if float(val).is_integer() else '#,##0.00'
    else:
        for col_idx, col_name in enumerate(pivot_df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=str(col_name))
            cell.fill = top_header_fill; cell.font = top_header_font
        for r_idx, (idx_val, row_data) in enumerate(pivot_df.iterrows(), start=2):
            for c_idx, (col_name, val) in enumerate(zip(pivot_df.columns, row_data), start=1):
                val_cell = ws.cell(row=r_idx, column=c_idx, value=val)
                if str(col_name).upper() in ['ACOS', 'TACOS'] and isinstance(val, (int, float, np.number)):
                    val_cell.value = val / 100.0 if val > 1 else val
                    val_cell.number_format = '0.00%'

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col[0].column)].width = max(max_len + 3, 10)
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def convert_all_pivots_to_excel(pivot_dict):
    buffer = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheets_added = 0
    for sheet_name, pivot_df in pivot_dict.items():
        if pivot_df is not None and not pivot_df.empty:
            temp_wb = openpyxl.load_workbook(io.BytesIO(style_and_export_pivot(pivot_df, sheet_name=sheet_name)))
            temp_ws = temp_wb.active
            temp_ws.title = sheet_name[:31]
            wb._sheets.append(temp_ws)
            sheets_added += 1
    if sheets_added == 0:
        wb.create_sheet(title="Summary").cell(row=1, column=1, value="No data available.")
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def convert_dataframe_to_excel_download(df_to_export, sheet_name="Sheet1"):
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        df_to_export.to_excel(writer, sheet_name=sheet_name[:31], index=False)
    buffer.seek(0)
    return buffer.getvalue()

st.markdown("### 📤 Upload Ad Campaign Reports (Up to 5)")
col1, col2, col3, col4, col5 = st.columns(5)
uploaded_files = []
file_types = ["csv", "xlsx", "xls", "xlsb", "xlsm"]
with col1:
    f1 = st.file_uploader("Upload Ad File 1", type=file_types, key="file1")
    if f1: uploaded_files.append(f1)
with col2:
    f2 = st.file_uploader("Upload Ad File 2", type=file_types, key="file2")
    if f2: uploaded_files.append(f2)
with col3:
    f3 = st.file_uploader("Upload Ad File 3", type=file_types, key="file3")
    if f3: uploaded_files.append(f3)
with col4:
    f4 = st.file_uploader("Upload Ad File 4", type=file_types, key="file4")
    if f4: uploaded_files.append(f4)
with col5:
    f5 = st.file_uploader("Upload Ad File 5", type=file_types, key="file5")
    if f5: uploaded_files.append(f5)

st.markdown("### 📥 Upload Sales Reports (Up to 5)")
scol1, scol2, scol3, scol4, scol5 = st.columns(5)
uploaded_sales_files = []
with scol1:
    s1 = st.file_uploader("Upload Sales File 1", type=file_types, key="sales1")
    if s1: uploaded_sales_files.append(s1)
with scol2:
    s2 = st.file_uploader("Upload Sales File 2", type=file_types, key="sales2")
    if s2: uploaded_sales_files.append(s2)
with scol3:
    s3 = st.file_uploader("Upload Sales File 3", type=file_types, key="sales3")
    if s3: uploaded_sales_files.append(s3)
with scol4:
    s4 = st.file_uploader("Upload Sales File 4", type=file_types, key="sales4")
    if s4: uploaded_sales_files.append(s4)
with scol5:
    s5 = st.file_uploader("Upload Sales File 5", type=file_types, key="sales5")
    if s5: uploaded_sales_files.append(s5)

sales_dfs_list = []
if uploaded_sales_files:
    for s_file in uploaded_sales_files:
        try:
            if s_file.name.lower().endswith('.csv'):
                sales_dfs_list.append(pd.read_csv(s_file, low_memory=False))
            else:
                s_xls = pd.ExcelFile(s_file)
                for s_sheet in s_xls.sheet_names:
                    sales_dfs_list.append(pd.read_excel(s_xls, sheet_name=s_sheet))
        except Exception as e:
            st.error(f"Error reading sales file {s_file.name}: {str(e)}")

master_sales_df = pd.DataFrame()
if sales_dfs_list:
    raw_sales_concat = pd.concat(sales_dfs_list, ignore_index=True)
    prod_col = next((c for c in ['Product Name', 'Item Name', 'product_name', 'item_name', 'Product'] if c in raw_sales_concat.columns), None)
    date_col = next((c for c in ['Order Date', 'date', 'Date', 'ORDER_DATE', 'timestamp'] if c in raw_sales_concat.columns), None)
    qty_col = next((c for c in ['Quantity', 'Qty', 'Units', 'qty_sold'] if c in raw_sales_concat.columns), None)
    mrp_col = next((c for c in ['MRP (Rs)', 'mrp', 'MRP', 'Selling Price (Rs)'] if c in raw_sales_concat.columns), None)
    city_col = next((c for c in ['Supply City', 'city_name', 'City'] if c in raw_sales_concat.columns), None)
    state_col = next((c for c in ['Supply State', 'Region', 'State'] if c in raw_sales_concat.columns), None)
    
    master_sales_df = raw_sales_concat.copy()
    master_sales_df['_product_name'] = master_sales_df[prod_col].astype(str) if prod_col else ""
    master_sales_df = master_sales_df[master_sales_df['_product_name'].str.lower().str.contains('staze', na=False)]
    
    if date_col:
        master_sales_df['_sales_dt'] = pd.to_datetime(master_sales_df[date_col], errors='coerce')
        master_sales_df['Month'] = master_sales_df['_sales_dt'].dt.strftime('%B').str.upper()
        master_sales_df['Date_str'] = master_sales_df['_sales_dt'].dt.strftime('%Y-%m-%d')
        master_sales_df['Day'] = master_sales_df['_sales_dt'].dt.day
    else:
        master_sales_df['Month'] = "UNKNOWN"
        master_sales_df['Date_str'] = np.nan
        master_sales_df['Day'] = 1
        
    master_sales_df['_qty'] = pd.to_numeric(master_sales_df[qty_col].astype(str).str.replace(',', ''), errors='coerce').fillna(1.0) if qty_col else 1.0
    master_sales_df['_gross_amt'] = pd.to_numeric(master_sales_df[mrp_col].astype(str).str.replace(',', ''), errors='coerce').fillna(0.0) if mrp_col else 0.0
    master_sales_df['Supply City'] = master_sales_df[city_col] if city_col else "Unknown"
    master_sales_df['Supply State'] = master_sales_df[state_col] if state_col else "Unknown"

if uploaded_files:
    consolidated_dfs, raw_files_dict, consolidated_raw_tabs = [], {}, {}
    for uploaded_file in uploaded_files:
        fallback_month = os.path.splitext(uploaded_file.name)[0].upper()
        raw_files_dict[uploaded_file.name] = {}
        try:
            if uploaded_file.name.lower().endswith('.csv'):
                df = pd.read_csv(uploaded_file, low_memory=False)
                raw_files_dict[uploaded_file.name]["Sheet1"] = df.copy()
                date_col = next((c for c in ['Date', 'date', 'Day', 'DATE', 'timestamp'] if c in df.columns), None)
                dt_series = pd.to_datetime(df[date_col], errors='coerce') if date_col else pd.Series(pd.NaT, index=df.index)
                df['Month'] = dt_series.dt.strftime('%B').str.upper().fillna(fallback_month)
                consolidated_raw_tabs.setdefault("Sheet1", []).append(df.copy())
                consolidated_dfs.append(df)
            else:
                xls = pd.ExcelFile(uploaded_file)
                for sheet_name in [s for s in xls.sheet_names if s.lower() not in ['summary', 'consolidated_master']]:
                    df = pd.read_excel(xls, sheet_name=sheet_name)
                    raw_files_dict[uploaded_file.name][sheet_name] = df.copy()
                    date_col = next((c for c in ['Date', 'date', 'Day', 'DATE', 'timestamp'] if c in df.columns), None)
                    dt_series = pd.to_datetime(df[date_col], errors='coerce') if date_col else pd.Series(pd.NaT, index=df.index)
                    df['Month'] = dt_series.dt.strftime('%B').str.upper().fillna(fallback_month)
                    consolidated_raw_tabs.setdefault(sheet_name, []).append(df.copy())
                    consolidated_dfs.append(df)
        except Exception as e:
            st.error(f"Error processing {uploaded_file.name}: {str(e)}")

    if consolidated_dfs:
        final_df = pd.concat(consolidated_dfs, ignore_index=True)
        camp_col = next((c for c in ['Campaign Name', 'Campaign', 'CampaignId'] if c in final_df.columns), None)
        if camp_col: final_df = final_df[final_df[camp_col].astype(str).str.lower().str.contains('staze', na=False)]

        def get_numeric(cols):
            for c in cols:
                if c in final_df.columns: return pd.to_numeric(final_df[c].astype(str).str.replace(',', ''), errors='coerce').fillna(0.0)
            return pd.Series(0.0, index=final_df.index)

        final_df['_impressions'] = get_numeric(['Impressions'])
        final_df['_atc'] = get_numeric(['Direct ATC']) + get_numeric(['Indirect ATC'])
        final_df['_orders'] = get_numeric(['Direct Quantities Sold', 'Direct Orders']) + get_numeric(['Indirect Quantities Sold', 'Indirect Orders'])
        final_df['_sales'] = get_numeric(['Direct Sales']) + get_numeric(['Indirect Sales'])
        final_df['_budget_consumed'] = get_numeric(['Estimated Budget Consumed', 'Budget Consumed', 'Spend'])
        final_df['_most_viewed_position'] = get_numeric(['Most Viewed Position', 'Average Position'])
        final_df['_date_dt'] = pd.to_datetime(final_df['Date'], errors='coerce') if 'Date' in final_df.columns else pd.NaT

        def create_monthly_summary(df_in):
            working = df_in.dropna(subset=['Month']).copy()
            if working.empty: return pd.DataFrame()
            rows = []
            for m in get_calendar_months_asc(working):
                m_df = working[working['Month'] == m]
                imp, atc, ords, spends, sales = m_df['_impressions'].sum(), m_df['_atc'].sum(), m_df['_orders'].sum(), m_df['_budget_consumed'].sum(), m_df['_sales'].sum()
                d = {'Month': m, 'Impressions': imp, 'CPM': round((spends/imp)*1000, 2) if imp>0 else 0.0, 'ATC': atc, 'Spends': spends, 'Ad Sales': sales}
                if not master_sales_df.empty:
                    m_sales = master_sales_df[master_sales_df['Month'] == m]
                    ov_sales = m_sales['_gross_amt'].sum()
                    d.update({'Overall Sales': ov_sales, 'Ad ROAS': round(sales/spends, 2) if spends>0 else 0.0, 'TROAS': round(ov_sales/spends, 2) if spends>0 else 0.0, 'TACOS': f"{round((spends/ov_sales)*100, 2) if ov_sales>0 else 0.0:.2f}%"})
                d['ACOS'] = f"{round((spends/sales)*100, 2) if sales>0 else 0.0:.2f}%"
                rows.append(d)
            return pd.DataFrame(rows)

        st.markdown("### 🔍 Global Dashboard Filter")
        sel_month = st.selectbox("Select Month:", ["All Months"] + get_calendar_months_asc(final_df))
        filt_df = final_df if sel_month == "All Months" else final_df[final_df['Month'] == sel_month]

        tot_sales, tot_spends = round(filt_df['_sales'].sum(), 2), round(filt_df['_budget_consumed'].sum(), 2)
        m1, m2, m3 = st.columns(3)
        with m1: st.metric("Ad Sales", f"₹{int(round(tot_sales)):,}")
        with m2: st.metric("Spends", f"₹{int(round(tot_spends)):,}")
        with m3: st.metric("Ad ROAS", f"{round(tot_sales/tot_spends, 2) if tot_spends>0 else 0.0:.2f}x")

        st.divider()
        tab1, tab2, tab3 = st.tabs(["📄 Raw Preview", "📊 Monthly Summary MoM", "💾 Master Dumps & Reports"])
        
        with tab1:
            st.dataframe(final_df.head(100), use_container_width=True)
            
        with tab2:
            ms_df = create_monthly_summary(final_df)
            st.dataframe(ms_df, use_container_width=True)
            st.download_button("📥 Download Monthly Summary (.xlsx)", style_and_export_pivot(ms_df, "Monthly_Summary"), "Monthly_Summary.xlsx")

        with tab3:
            st.markdown("#### Download Master Workbook with ADs Dump & Overall Dump")
            ads_dump = final_df.copy()
            ads_dump['date_ist'] = ads_dump['_date_dt'].dt.strftime('%Y/%m/%d') if '_date_dt' in ads_dump.columns else np.nan
            
            final_ads_export = pd.DataFrame({
                'Campaign Name': ads_dump.get('Campaign Name', ads_dump.get('Campaign', np.nan)),
                'Impressions': ads_dump['_impressions'],
                'Sales': ads_dump['_sales'],
                'Spends': ads_dump['_budget_consumed'],
                'date_ist': ads_dump['date_ist']
            })

            final_overall_export = pd.DataFrame({
                'item_name': master_sales_df.get('_product_name', np.nan),
                'city_name': master_sales_df.get('Supply City', np.nan),
                'qty_sold': master_sales_df.get('_qty', np.nan),
                'mrp': master_sales_df.get('_gross_amt', np.nan)
            }) if not master_sales_df.empty else pd.DataFrame(columns=['item_name', 'city_name', 'qty_sold', 'mrp'])

            buf = io.BytesIO()
            with pd.ExcelWriter(buf, engine='openpyxl') as writer:
                final_ads_export.to_excel(writer, sheet_name='ADs Dump', index=False)
                final_overall_export.to_excel(writer, sheet_name='Overall Dump', index=False)
                for r_name, d_list in consolidated_raw_tabs.items():
                    pd.concat(d_list, ignore_index=True).to_excel(writer, sheet_name=r_name[:31], index=False)
            buf.seek(0)

            st.download_button(
                label="📥 Download Complete Excel Workbook (ADs Dump + Overall Dump + Raw Tabs)",
                data=buf,
                file_name="STAZE_Blinkit_Master_Report.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary"
            )
    else:
        st.info("Please upload valid ad campaign reports.")
else:
    st.info("Awaiting file uploads.")
